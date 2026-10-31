"""Reproduce the Chinese operator manual with python-docx 1.2.0.

Run in the document tool container. Screenshots in .runtime/manual-qa/images
are optional; final rendering must be checked after regeneration.
"""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'manual' / 'LicenseServer使用說明書.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.header_distance = sec.footer_distance = Cm(.8)
for name, size in [('Normal',10),('Title',26),('Subtitle',13),('Heading 1',18),('Heading 2',12),('Heading 3',11)]:
    style = doc.styles[name]
    style.font.name = 'Noto Sans CJK TC'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK TC')
    style.paragraph_format.space_after = Pt(7 if name == 'Normal' else 10)
    style.paragraph_format.line_spacing = 1.18
doc.styles['Normal'].paragraph_format.widow_control = True
header = sec.header.paragraphs[0]
header.text = 'LicenseServer 8443  使用說明書'
header.style = doc.styles['Normal']
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string('536578')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026 10 31  ｜  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
for r in footer.runs: r.font.size = Pt(8)

def p(text, bold=False):
    para = doc.add_paragraph()
    r = para.add_run(text); r.bold = bold
    return para

def h(text): doc.add_heading(text,2)
def page(title):
    doc.add_page_break(); doc.add_heading(title,1)
def steps(items):
    for i,text in enumerate(items,1): p(f'{i}　{text}')

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for col,w in zip(t.columns,widths): col.width = Cm(w)
    for i,text in enumerate(headers): t.rows[0].cells[i].text = text
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for i,text in enumerate(row): cells[i].text = str(text)
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for ci,cell in enumerate(row.cells):
            if widths: cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'173B59' if ri==0 else ('F2F5F8' if ri%2 else 'FFFFFF')); props.append(shade)
            margins=OxmlElement('w:tcMar')
            for side,val in [('top',85),('bottom',85),('left',95),('right',95)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1)
                para.paragraph_format.line_spacing=1.1
                for r in para.runs:
                    r.font.size=Pt(9)
                    if ri==0: r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def screenshot(name, caption, width=16.7):
    path=ROOT/'.runtime/manual-qa/images'/name
    if path.exists():
        doc.add_picture(str(path),width=Cm(width))
        para=p(caption);para.paragraph_format.space_after=Pt(5)
        for r in para.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('536578')

doc.add_paragraph('LicenseServer 使用說明書', 'Title')
doc.add_paragraph('8443 管理介面與 SBOMScan 授權操作', 'Subtitle')
p('本手冊提供核發與更新 SBOMScan 客戶授權的實際操作順序，並盤點目前 8443 管理介面的功能。日常操作採基本模式；目前 SBOMScan 未使用的管理入口收在「設定 → 一般 → 進階功能」。')
p('適用版本為本機 LicenseServer 與目前 SBOMScan 原始碼，盤點日期為 2026 年 10 月 31 日。功能有無入口、是否受版本限制、是否由 SBOMScan 使用，分開說明。進階模式只控制介面顯示，不會停用已存在的 API、Webhook 或變更授權內容。',True)
h('先找到你要做的事')
table(['工作','操作位置','手冊頁次'],[
('第一次核發客戶授權','使用者 → 授權 → 設備 → 簽出設備','6 至 8'),
('更新功能或續約','授權 → 編輯／功能權限／續約','7'),
('讓客戶離線使用','設備 → 簽出設備 → SBOMScan 匯入','8'),
('查看所有功能','基本功能盤點與進階功能盤點','2 至 4'),
('顯示技術管理入口','設定 → 一般 → 啟用進階功能','9'),
('匯入後沒有變化或不能上傳','故障排除與部署檢查','11 至 12'),
],[4,9,4.4])
h('日常操作的關係')
p('產品 SBOMScan → 授權原則提供方案預設 → 授權代表客戶的使用權 → 設備綁定實際部署 → Machine File 帶回該部署驗證。功能權限控制能做什麼，額度控制能使用多少。')
p('管理網址：https://localhost:8443。若從其他電腦連線，需由維運人員先配置可達網址、連接埠及受信任憑證；目前 Compose 的 8443 綁定在主機的 127.0.0.1。')

page('1 基本功能盤點')
p('基本模式保留核發、設備綁定與管理員自助設定。清單通常提供搜尋、篩選、排序、分頁與收藏；各頁實際顯示的操作仍受登入角色、API 權限及資料狀態限制。')
table(['入口與路徑','主要操作','SBOMScan 的用途'],[
('首頁 dashboard','使用狀態摘要、到期情況、連到相關清單','看整體授權情況；此 CE 安裝不顯示未提供的進階分析區塊'),
('產品 products','新增、查看、編輯、刪除；連到授權原則、授權、設備與使用者','保留現有 SBOMScan 產品，供客戶端 Product ID 比對'),
('功能權限 entitlements','建立代碼與名稱、編輯、刪除；查看關聯','管理四個 SBOMSCAN 功能代碼；建立代碼後仍需附加給授權'),
('授權原則 policies','建立方案、複製、編輯、刪除；設定期限、功能與用量預設','提供核發時的方案預設；不要以為改原則就已更新所有舊授權'),
('授權 licenses','新增、複製、編輯、刪除、續約、停權／恢復、管理功能與使用者','每張客戶授權的到期、功能、設備數、同時使用人數及專案額度'),
('設備 machines','新增、查看、編輯、停用、重設心跳、簽出設備','用目標部署的指紋綁定授權；產生可匯入的 Machine File'),
('使用者 users','新增、編輯、刪除、指派授權、查看關聯；可用的角色操作依權限顯示','客戶或授權擁有者；與 SBOMScan 登入帳號分開管理'),
('設定 一般 general','個人／帳戶資料、語言、外觀、進階功能開關','日常設定與顯示簡化'),
('設定 安全性 security','密碼與雙因素驗證等帳號安全操作、本機 HTTPS 憑證下載','管理員登入安全與下載公開 CA'),
],[4,7,6.4])
p('刪除、停用與停權會改變資料或授權狀態。執行前應確認客戶、License ID 與 Machine ID；本手冊的介面盤點沒有執行這些破壞性操作。')

page('2 進階功能盤點')
p('下列入口在基本模式隱藏。SBOMScan 目前不依賴這些資源完成核發、設備啟用及離線驗證；未來若另外串接，可在進階模式使用。這些功能與 SBOM 掃描結果中的套件或元件清單並非同一套資料。')
table(['入口與路徑','功能','目前為何收起'],[
('群組 groups','集中關聯授權、設備或使用者；管理群組擁有者與限制','目前授權驗證使用單張授權與設備指紋，不依群組判斷'),
('元件 components','管理設備下可分別識別的授權元件及指紋','SBOMScan 未建立 Keygen Components；不是 SBOM 裡的軟體元件'),
('程序 processes','管理設備的執行程序、租約與心跳','SBOMScan 的同時使用人數由產品端管理，不以 Keygen Process 紀錄代表登入人數'),
('套件 packages','組織可發布的軟體套件及其版本','目前由 Docker 與既有部署流程交付產品'),
('版本 releases','建立、編輯、發布或撤下版本；管理關聯產物','SBOMScan 未使用 Keygen Release 作為更新來源'),
('產物 artifacts','管理發布檔案、上傳或下載及關聯資訊','不是上傳 SBOM 專案的入口'),
('平台 platforms','管理發佈平台分類','目前產品未使用 Keygen 軟體發佈'),
('架構 arches','管理 CPU 架構分類','用於發佈檔案分類，並非設備授權數'),
('通道 channels','管理版本通道，例如 stable 或 beta','目前無 Keygen 發布通道串接'),
('引擎 engines','管理軟體發佈的引擎或生態分類','目前未使用此發佈分類'),
],[4,6.5,6.9])
h('基本頁面同步收起的入口')
p('授權、設備及使用者表單與詳細頁的群組入口收起；設備的元件與程序連結收起。設備簽出在基本模式只列出「授權」與「功能權限」；其他關聯選項在進階模式恢復。進階資料並未被刪除。')

page('3 技術與版本限制功能盤點')
table(['入口與路徑','功能與現況','顯示方式'],[
('Webhook 端點 webhook-endpoints','設定外部接收網址、事件範圍及相關交付設定；授權事件發生時通知其他系統','進階；目前 SBOMScan 依定期驗證或檔案匯入更新，不依賴 Webhook'),
('Webhook 事件 webhook-events','查看事件與交付情形；可用的重送等操作依 API 與事件狀態提供','進階；設定端點不等於已串好 SBOMScan'),
('Token tokens','建立、查看、撤銷或管理 API 憑證；角色及範圍依 API 支援','進階；供供應商維運或程式串接，不能當成客戶 License Key'),
('事件記錄 event-logs','原生稽核事件清單及詳細資料','進階；此安裝顯示不可用提示，開啟進階不會取得額外版本能力'),
('請求記錄 request-logs','原生 API 請求記錄與詳細資料','進階；此安裝不提供完整請求記錄'),
('團隊 team','管理登入授權伺服器的管理團隊與角色','進階；不是 SBOMScan 專案成員管理'),
('權限 permissions','帳戶層級的細部 API 權限設定','進階分類；原入口有 EE 限制，CE 不會因此解鎖'),
('開發人員 developers','帳戶整合設定、簽章用公私鑰及開發用 Token 區塊','進階；公鑰可用於驗證，私鑰及管理 Token 只留在供應商端'),
('帳務 billing','代管服務帳務管理','進階分類；目前自架 CE 不適用，既有版本隱藏規則保留'),
('學習 learn','Keygen 入門精靈、建立產品或原則的教學','進階；目前環境已有 SBOMScan 產品與方案，避免重複建立'),
],[4,8.2,5.2])
p('以上與第 2 至 3 頁共涵蓋側邊欄來源的 29 個頁面入口；首頁的「指標」已整併到首頁按鈕。詳細頁、建立與編輯視窗屬於其對應資源，不另算一種授權功能。')
h('其他進階操作')
p('授權詳細頁的「簽出授權」產生通用 License File，與 SBOMScan 所需的 Machine File 不同，已收進進階模式。「重設用量」用於計次型授權，也收進進階模式。授權與原則表單內未映射的資源容量、計次、元件與部分驗證範圍欄位亦收起；維運設定請參照第 10 頁。')

page('4 名詞與 SBOMScan 對應')
table(['名稱','意思與設定重點'],[
('產品 Product','授權對象，例如 SBOMScan。客戶端要使用同一個 Account ID、Product ID 與正確驗證公鑰。'),
('授權原則 Policy','方案的共用預設規則；單張授權可能覆寫期限或額度。現有方案包括 Trial、Basic、Professional、Enterprise。'),
('授權 License','核給特定客戶的權利。License Key 用於啟用；License ID 用於管理與追蹤。'),
('設備 Machine','一套實際部署的啟用紀錄。設備指紋由 SBOMScan 產生，不能自行編一段名稱代替。'),
('使用者與擁有者','LicenseServer 的客戶資料與授權歸屬。SBOMScan 的 admin、analyst 與專案成員由產品自己的帳號系統處理。'),
('功能權限 Entitlement','字串代碼控制產品功能。顯示名稱可讀，真正比對以代碼為準。'),
],[4,13.4])
h('目前支援的功能代碼')
table(['代碼','產品端用途'],[
('SBOMSCAN_SCAN','掃描相關操作與功能入口'),
('SBOMSCAN_COLLECTOR','Collector 資料收集流程'),
('SBOMSCAN_REMEDIATION','修復管理'),
('SBOMSCAN_VULNDB_UPDATE','漏洞資料庫更新'),
],[9,8.4])
h('三種額度要分開')
p('設備數 maxMachines 限制同一張授權可啟用的部署數；同時使用人數由 maxProcesses 與 metadata.concurrent_users 對應，更新時兩者維持相同正整數；每套部署的專案額度使用 metadata.max_projects，正整數表示上限，-1 表示不限。')
p('maxUsers 並非 SBOMScan 的同時登入人數。若 concurrent_users 與 maxProcesses 不一致，產品端可能優先讀取中繼資料；只改其中一個可能看不到預期結果。功能權限、額度及使用者角色需要同時符合，才可使用對應操作。')

page('5 核發新的客戶授權')
steps([
'登入 8443，先到「產品」確認現有 SBOMScan，再到「授權原則」選擇適合的方案。一般核發不需要重新建立產品。',
'到「使用者」搜尋客戶 Email；尚未存在再新增。這筆資料用於客戶識別與授權擁有者，不會自動建立 SBOMScan 登入帳號。',
'到「授權」按新增，選取產品對應的原則與客戶擁有者，填入易辨識的客戶名稱及到期時間。日期需確認時區。',
'設定設備數、同時使用人數與每套部署的專案額度。專案額度為必填，Trial 預設 3、Basic 預設 5，其他方案由核發者指定；如需不限，使用明確的不限選項。',
'附加應開放的功能權限，核發後再次查看該張授權的功能清單。要使用掃描，必須具備 SBOMSCAN_SCAN；只建立 Entitlement 或只改名稱並不會開權限。',
'儲存並記下 License ID、客戶、到期日與核定額度。線上／首次啟用客戶取得 License Key；完全離線客戶需要依第 8 頁建立對應設備並簽出 Machine File。',
])
h('核發前核對表')
table(['欄位','應確認內容'],[
('產品與原則','屬於現有 SBOMScan；避免選到別的產品'),
('到期與狀態','期限符合合約；授權未停權且已到生效時間'),
('功能權限','以四個精確代碼判斷；額度增加不會自動增加功能'),
('同時使用人數','maxProcesses 與 metadata.concurrent_users 一致，皆為正整數'),
('專案額度','metadata.max_projects 為正整數或 -1，不能用 0 或字串表示不限'),
('設備上限','涵蓋客戶要啟用的部署數；舊設備是否仍占用名額'),
],[4,13.4])
p('產品預設與原則的細部限制可能影響核發結果，核發完成後應以該張 License 的實際值為準。修改原則中的專案預設不會自動回寫舊 License；已有客戶需要逐張更新。')

page('6 更新功能 額度與到期日')
steps([
'到「授權」找到原客戶的 License，以 License ID 確認，而不是新增一張相似名稱的授權。',
'更新該張授權的功能權限，或編輯專案額度、設備數與同時使用人數。若要延長期限，使用續約或編輯到期日；確認操作結果的實際日期。',
'同時使用人數要同步檢查 maxProcesses 與 metadata.concurrent_users。原本中繼資料有值時，只改「程序上限」可能仍沿用舊人數。',
'先儲存，再確認 License 詳細頁的新值。使用離線檔或 activation 模式的客戶，依第 8 頁在同一個 Machine 重新簽出，取得新檔。',
'客戶於 SBOMScan「設定 → 產品授權」匯入新 Machine File；若為 online，確認客戶端成功連線並更新授權狀態。',
'查看產品授權頁的狀態、到期、功能及額度，再實際開啟對應功能確認。未授權功能的選單會隱藏；新增授權後重新讀取頁面以取得新入口。',
])
h('依授權模式判斷更新方式')
table(['模式','如何取得更新','需要注意'],[
('online','保持連線，等待週期驗證或由授權頁重新啟用／更新','預設約每分鐘驗證；斷線期間不會即時收到中央變更'),
('activation','首次用 License Key 連線啟用，之後主要驗證本機檔；更新時重新連線啟用或匯入新檔','中央改完後，不要假設本機檔已自動換新'),
('offline','在 8443 重新簽出正確 Machine File，再到產品授權頁匯入','舊檔不會因伺服器資料更新而改變'),
('disabled','開發測試模式，不作正式授權驗證','不能用此模式驗收客戶權限；正式版應拒絕使用 disabled'),
],[3,8.4,6])
p('中央停權或調降額度不會遠端改寫已交付且完全離線的舊檔。離線客戶的變更需靠新檔、生效／到期時間與客戶端重新匯入流程落實。若調降專案額度至低於既有數量，保留既有資料，但禁止繼續新增。')

page('7 設備綁定與正確簽出')
p('SBOMScan 匯入的是設備簽章檔，檔案開頭為 -----BEGIN MACHINE FILE-----。請從「設備」簽出，不是從「授權」簽出通用 License File。')
steps([
'在客戶目標 SBOMScan 的「設定 → 產品授權」取得設備指紋。每套部署各自使用持久化的安裝識別與授權資料，不能複製別套部署的授權 volume 代替。',
'到 8443「設備」找到該部署。線上啟用通常已建立 Machine；若為完全離線流程，新增設備、選對客戶 License，並貼上目標部署的精確指紋。',
'進入該台設備詳細頁，從操作選單點「簽出設備」。核對 Machine ID 與所屬 License，勿選到另一台設備。',
'關聯資料至少保留「授權」及「功能權限」，目前已預先勾選。需要顯示方案名稱等資訊時，可在進階模式額外包含原則；不要移除必要的兩項。',
'使用 Ed25519，關閉加密，設定有限且符合需求的有效期。預設檔案期限為 30 天；License 的到期仍會另外檢查。其他算法或加密格式不符合目前產品匯入器。',
'完成簽出、下載或複製完整檔案，於原目標 SBOMScan「設定 → 產品授權」匯入。確認匯入成功、狀態有效、設備指紋一致及功能額度正確。',
])
h('期限與設備特性')
p('有效授權同時受檔案期限、License 到期、生效時間、停權狀態及設備綁定影響。延長 License 期限後，舊 Machine File 的期限不會自動延長，必須重新簽出。簽出檔案的 TTL 應至少符合 API 的最小限制，接近到期時不要以不足一小時的期限嘗試簽出。')
p('重建客戶環境時需保留自己的 product_license volume 與安裝識別。換設備或重建識別可能產生新指紋，需重新綁定；停用舊設備前先確認遷移安排與上限。停用伺服器上的設備也不代表已交付的完全離線檔立刻失效。')

page('8 開啟或關閉進階功能')
steps([
'點左側「設定」，進入「一般」。',
'找到「進階功能」，點「啟用進階功能」。介面會重新載入，顯示群組、元件、程序、發布、Webhook、Token 與技術管理入口。',
'完成操作後回到相同位置，點「關閉進階功能」，恢復日常簡化介面。',
])
p('偏好儲存在目前瀏覽器的本機儲存空間。換瀏覽器、使用無痕模式或清除網站資料後，可能回到基本模式。若無法保存，請確認瀏覽器允許此網站使用儲存空間。')
screenshot('settings.png','圖 1　一般設定中的進階功能開關與目前模式')
h('這個開關會影響哪些位置')
p('側邊欄與功能類別、相關收藏、快速操作指令及最近使用的進階資源，會依模式調整。直接開啟被收起的進階頁面，會顯示提示並提供回到設定的入口。基本模式仍可進入帳號安全頁、下載公開 HTTPS CA 以及管理日常授權。')
p('進階功能不是帳號權限，也不是 SBOMScan 的功能授權旗標。開啟後仍須具備原 API 權限；CE 不提供的功能仍受版本限制。關閉後不會刪除既有資料，也不會停用已設定的事件交付。')

page('9 進階設定操作參考')
p('進階管理由維運或整合人員使用。變更前先記錄現值與受影響的客戶，尤其是共用授權原則；可先以獨立測試授權驗證。')
table(['功能或設定','用途與使用原則'],[
('授權原則的期限策略','duration、到期基準與續約基準控制到期計算。操作後以 License 實際日期驗證，不能只看方案名稱。'),
('設備唯一性與驗證範圍','strict、機器匹配及 machine／fingerprint scope 影響綁定檢查；應維持 SBOMScan 相容設定。元件、checksum、version、user scope 需客戶端提供相應資料，不能隨意要求。'),
('心跳與租約','Keygen 設備／程序心跳、租約回收與復活策略。SBOMScan 的定期授權驗證不等於已實作這些心跳協定。不要任意啟用 requireHeartbeat 或 requireCheckIn。'),
('資源與計次用量','maxCores、maxMemory、maxDisk、maxUses、usePool、overage 等為原生用量模型。SBOMScan 目前主要映射設備數、同時使用人數及專案數，不會自動把其他欄位當成產品配額。'),
('簽出授權與重設用量','前者產生通用 License File；後者重設 Keygen 計次用量。兩者都不取代 SBOMScan 的設備檔匯入，也不會清除產品專案數。'),
('Webhook','新增端點、選擇事件、由接收端驗證並處理通知，再觀察交付紀錄。需另有接收服務，只有新增 URL 不會完成產品權限同步。'),
('API Token 與開發人員','供應商程式串接用。給客戶的是 License Key、Machine File 與必要的公開驗證設定；不要交付管理員 Token、私鑰或授權資料庫。'),
('團隊與使用者權限','團隊控制授權伺服器的管理工作；SBOMScan 使用者角色與專案成員需在產品內設定。'),
],[5,12.4])
p('進階模式可編輯被收起的核心數、記憶體、磁碟、使用次數、Keygen 使用者數、用量池、元件驗證與程序租約等欄位。基本模式仍可檢視詳細頁原生技術屬性，以核對歷史設定；切換模式不會重設既有值。心跳等可能影響現有授權的設定保留於原生表單供維運確認。')

page('10 匯入與功能異常排除')
table(['現象','優先檢查與處理'],[
('匯入後看不到新功能','先看產品授權頁是否成功且有效。核對同一 License／Machine；新檔是否在修改儲存後才簽出；包含 license 及 license.entitlements；新功能代碼是否真的附加；重新載入產品頁。'),
('顯示設備或產品不符','比較指紋、Account ID、Product ID 與公鑰。用目標部署建立的 Machine 重新簽出，不能把其他部署檔案拿來匯入。'),
('無法建立或上傳專案','核對 SBOMSCAN_SCAN、產品端角色權限、metadata.max_projects 與已使用量。舊檔缺專案額度會禁止建立新專案；更新 License 額度並重新簽出匯入。額度足夠後才查上傳格式、大小與伺服器日誌。'),
('提高人數卻沒有變化','核對 maxProcesses 與 metadata.concurrent_users 是否同步；確認客戶已取得新檔或線上更新成功；避免只改 maxUsers。'),
('簽出一直失敗','看畫面錯誤提示與 API 日誌；檢查設備及授權仍有效、TTL 合法、剩餘期限是否不足 API 最小值。避免選擇目前匯入器不支援的加密或簽章算法。'),
('期限還沒到卻是唯讀','同時檢查檔案期限、License 期限、生效時間、系統時間、停權及線上斷線寬限。License 延長不會更新舊簽章檔。'),
('授權有效但選單仍不顯示','產品同時檢查功能授權與使用者角色。未授權入口會隱藏；有權限卻到期轉唯讀時，依產品設計保留歷史資料檢視。'),
('找不到元件或 Webhook','到 8443「設定 → 一般」啟用進階功能。這與 SBOMScan 的功能開關無關。'),
('看到事件記錄不可用','目前 CE 安裝未提供該能力；進階開關不會解鎖。需要查故障可先讀取 Docker 服務日誌。'),
],[5,12.4])
p('回報問題時提供發生時間、License ID、Machine ID、授權模式、畫面錯誤與已遮蔽秘密的日誌。不要在問題截圖或公開 GitHub issue 放完整 License Key、Machine File、Token 或私鑰。')

page('11 部署檢查與盤點依據')
h('啟動與查看狀態')
p('在 LicenseServer 資料夾執行 docker compose up -d。查看 docker compose ps -a 與 docker compose logs --tail 100 portal web worker gateway；setup、catalog、runtime-permissions 屬於一次性工作，正常結束為 Exited 0。')
p('停止並保留資料可執行 docker compose down。不要加 -v，因為這會移除資料 volume。授權伺服器的資料庫、設定及 .runtime 應依維運備份流程保留。')
h('客戶端重新部署後的驗收')
steps([
'使用正式授權映像，確認 LICENSE_MODE 是 online、activation 或 offline。只刪除「後門」資料夾不會改變映像或既有容器的模式。',
'確認客戶自己的 product_license volume 與安裝識別仍保留；本機測試連線需使用可從容器到達的授權伺服器網址與正確 CA。',
'登入 SBOMScan「設定 → 產品授權」，確認有效狀態、功能、到期及專案額度；用一般使用者檢查角色限制，並驗證未授權入口會隱藏。',
'實際測試一次核准範圍內的專案建立／上傳與掃描。只看容器啟動成功，不代表授權與產品流程都已成功。',
])
h('本次盤點範圍與依據')
p('盤點包含目前管理介面的側邊欄資源、詳細頁主要操作、建立／編輯表單、簽出流程、一般／安全設定，以及 CE 的既有顯示限制。功能名稱與路徑來自固定版本的 Portal，授權對應來自本機 SBOMScan 驗證與配額程式。')
table(['來源','核對內容'],[
('Portal 固定版本','keygen-sh/keygen-portal 提交 2d4bb657273528aafdfd725cd7203021a067c742；src/lib/views.ts、src/pages/app、各資源 form 與 schemas'),
('LicenseServer 客製程式','portal-customize.mjs、portal-advanced.mjs、portal-i18n/advanced-features.ts 與 tsx、product.py、compose.yaml'),
('SBOMScan 客戶端','LICENSING.md；server/internal/licensing/manager.go、certificate.go、projects.go 與相關驗證流程'),
('實際驗證','8443 基本與進階模式切換、主要清單及直接進階入口；建置時執行 TypeScript、翻譯與配額／顯示規則測試'),
],[4.3,13.1])
p('此手冊不代表每個進階資源已完成客戶端串接，也沒有新增測試客戶、簽發測試授權、變更正式額度或執行刪除。後續若更換 Portal 或 SBOMScan 版本，應重新核對畫面、對應代碼與簽章格式。')

doc.core_properties.title='LicenseServer 使用說明書'
doc.core_properties.subject='8443 功能盤點與 SBOMScan 客戶授權操作'
doc.core_properties.author='LicenseServer 維運文件'
doc.core_properties.keywords='LicenseServer SBOMScan 8443 授權 進階功能'
# Remove template title borders and theme font overrides explicitly.
for node in [doc.styles.element, doc.element]:
    for border in list(node.xpath('.//w:pBdr')):
        border.getparent().remove(border)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if key.endswith('Theme'): del fonts.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+key),'Noto Sans CJK TC')
    style.font.italic=False
    style.font.underline=False
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
